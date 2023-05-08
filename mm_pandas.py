
import pandas as pd
from docx import Document

# Read the Excel data
excel_file = 'data.xlsx'
data = pd.read_excel(excel_file, engine='openpyxl')

# Load the Word template
template_file = 'template.docx'

# Iterate over each row in the Excel data
for index, row in data.iterrows():
    # Create a new document from the template
    doc = Document(template_file)

    # Replace placeholders in the document with data from the Excel file
    for paragraph in doc.paragraphs:
        for key in row.keys():
            placeholder = f'<<{key}>>'
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, str(row[key]))

    # Save the new document
    output_file = f'output_{index + 1}.docx'
    doc.save(output_file)

print("Mail merge complete.")